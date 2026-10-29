
from django.core.serializers import python
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.shortcuts import render, redirect
from django.contrib.auth.models import User
from django.contrib.auth import authenticate, login, logout

from .models import OwnerProfile
from django.shortcuts import render
from django.shortcuts import render
from .models import HomeHero

from django.shortcuts import render, redirect, get_object_or_404
# from .models import HomeFeature
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from .models import HomeCategory
# Add with your other imports
from .models import Property, Review
from .models import WhyChooseFeature

# def home(request):
#     return render(request, 'home.html')

def home(request):
    why_choose_features = WhyChooseFeature.objects.filter(
        is_active=True
    ).order_by('display_order')

    context = {
        'why_choose_features': why_choose_features,
    }

    return render(request, 'home.html', context)

def admin_login(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None:

            if hasattr(user, 'adminprofile'):
                login(request, user)
                return redirect('admin_dashboard')

            else:
                return render(
                    request,
                    'admin_panel/admin_login.html',
                    {
                        'error': 'You are not authorized as Admin.'
                    }
                )

        else:
            return render(
                request,
                'admin_panel/admin_login.html',
                {
                    'error': 'Invalid username or password.'
                }
            )

    return render(
        request,
        'admin_panel/admin_login.html'
    )
   
# Admin Dashboard


from django.shortcuts import render, redirect
from django.contrib.auth.models import User
from .models import AdminProfile, OwnerProfile
from datetime import date, datetime

from django.contrib.auth.decorators import login_required
from django.shortcuts import redirect
from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from django.contrib.auth import get_user_model
from .models import AdminProfile, OwnerProfile, Property, Booking, Review

User = get_user_model()

def admin_dashboard(request):
    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not AdminProfile.objects.filter(user=request.user).exists():
        return redirect('admin_login')

    total_users = User.objects.filter(
        is_superuser=False,
        is_staff=False,
        ownerprofile__isnull=True
    ).count()

    total_owners = OwnerProfile.objects.count()

    report_users = User.objects.filter(
        is_superuser=False,
        is_staff=False,
        ownerprofile__isnull=True
    ).order_by('-date_joined')

    report_owners = User.objects.filter(
        ownerprofile__isnull=False
    ).order_by('-date_joined')

    report_properties = Property.objects.select_related(
        'owner'
    ).order_by('-created_at')

    report_bookings = Booking.objects.select_related(
        'user', 'property'
    ).order_by('-pk')

    report_reviews = Review.objects.select_related(
        'user', 'property'
    ).order_by('-created_at')

    context = {
        'total_users': total_users,
        'total_owners': total_owners,
        'total_properties': report_properties.count(),
        'total_bookings': report_bookings.count(),
        'total_reviews': report_reviews.count(),
        'report_users': report_users,
        'report_owners': report_owners,
        'report_properties': report_properties,
        'report_bookings': report_bookings,
        'report_reviews': report_reviews,
    }

    return render(
        request,
        'admin_panel/admin_dashboard.html',
        context
    )

def admin_logout(request):
    logout(request)
    return redirect('admin_login')

# Owner Profile Views



# =========================
# OWNER REGISTER
# =========================
def owner_register(request):

    if request.method == 'POST':

        full_name = request.POST.get('full_name')
        email = request.POST.get('email')
        phone = request.POST.get('phone')
        password = request.POST.get('password')
        confirm_password = request.POST.get('confirm_password')

        # Check password
        if password != confirm_password:
            return render(
                request,
                'owner_panel/owner_register.html',
                {'error': 'Passwords do not match.'}
            )

        # Check email
        if User.objects.filter(username=email).exists():
            return render(
                request,
                'owner_panel/owner_register.html',
                {'error': 'An account with this email already exists.'}
            )

        # Create User
        user = User.objects.create_user(
            username=email,
            email=email,
            password=password,
            first_name=full_name
        )

        # Create Owner Profile
        OwnerProfile.objects.create(
            user=user,
            phone=phone
        )

        return redirect('owner_login')

    return render(
        request,
        'owner_panel/owner_register.html'
    )


# =========================
# OWNER LOGIN
# =========================


from django.contrib.auth import authenticate, login, get_user_model
from django.shortcuts import render, redirect

User = get_user_model()


def owner_login(request):
    if request.method == 'POST':
        email_or_username = request.POST.get('email', '').strip()
        password = request.POST.get('password', '')

        # Find the user by email or username
        user_obj = User.objects.filter(
            email__iexact=email_or_username
        ).first()

        if user_obj is None:
            user_obj = User.objects.filter(
                username__iexact=email_or_username
            ).first()

        if user_obj is not None:
            user = authenticate(
                request,
                username=user_obj.get_username(),
                password=password
            )

            if user is not None:
                # Check whether the account belongs to an owner
                if hasattr(user, 'ownerprofile'):
                    if not user.ownerprofile.is_active:
                        return render(
                            request,
                            'owner_panel/owner_login.html',
                            {'error': 'Your owner account is inactive.'}
                        )

                    login(request, user)
                    return redirect('owner_dashboard')

                return render(
                    request,
                    'owner_panel/owner_login.html',
                    {'error': 'This account is not registered as an owner.'}
                )

        return render(
            request,
            'owner_panel/owner_login.html',
            {'error': 'Invalid email/username or password.'}
        )

    return render(request, 'owner_panel/owner_login.html')

# =========================
# OWNER DASHBOARD
# =========================

from django.shortcuts import render, redirect
from .models import Property, Booking, AvailableArea, Review


def owner_dashboard(request):

    # Check whether the owner is logged in
    if not request.user.is_authenticated:
        return redirect('owner_login')

    # Allow only accounts with an OwnerProfile
    if not hasattr(request.user, 'ownerprofile'):
        return redirect('owner_login')

    # Get properties belonging to the logged-in owner
    properties = Property.objects.filter(
        owner=request.user
    ).order_by('-id')

    # Get booking requests for this owner's properties
    owner_bookings = Booking.objects.filter(
        property__owner=request.user
    ).select_related(
        'property',
        'user'
    ).order_by('-id')

    # Get reviews for this owner's properties
    owner_reviews = Review.objects.filter(
        property__owner=request.user
    ).select_related(
        'property',
        'user'
    ).order_by('-id')

    # Get active areas added by Admin
    areas = AvailableArea.objects.filter(
        is_active=True
    ).order_by('display_order', 'name')

    # Available amenities for Add/Edit Property forms
    amenities_options = [
        'Wi-Fi',
        'Water Supply',
        'Power Backup',
        'Security',
        'Lift',
        'Washing Machine',
    ]

    # Dashboard context
    context = {
        'properties': properties,
        'total_properties': properties.count(),

        'booking_requests': owner_bookings.count(),
        'owner_bookings': owner_bookings,

        'owner_reviews': owner_reviews,
        'total_reviews': owner_reviews.count(),

        # Keep 'areas' for existing template code
        'areas': areas,

        # Use these names in the Add/Edit Property forms
        'available_areas': areas,
        'amenities_options': amenities_options,
    }

    return render(
        request,
        'owner_panel/owner_dashboard.html',
        context
    )

# ownwer report view
# =========================
# OWNER REPORTS
# =========================

from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .models import Property, Booking, Review


@login_required
def owner_reports(request):

    if not hasattr(request.user, 'ownerprofile'):
        return redirect('owner_login')

    # Get only this owner's data
    properties = Property.objects.filter(
        owner=request.user
    ).order_by('-id')

    bookings = Booking.objects.filter(
        property__owner=request.user
    ).select_related(
        'property', 'user'
    ).order_by('-id')

    reviews = Review.objects.filter(
        property__owner=request.user
    ).select_related(
        'property', 'user'
    ).order_by('-id')

    report_type = request.GET.get('report', 'properties')

    # Summary information
    context = {
        'report_type': report_type,
        'total_properties': properties.count(),
        'total_bookings': bookings.count(),
        'total_reviews': reviews.count(),
        'properties': properties,
        'bookings': bookings,
        'reviews': reviews,
    }

    return render(
        request,
        'owner_panel/owner_reports.html',
        context
    )
from datetime import date, datetime, time

from django.contrib.auth.decorators import login_required
from django.shortcuts import redirect
from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# =========================
# EXPORT REPORT TO EXCEL
# =========================

@login_required
def export_owner_reports(request):

    if not hasattr(request.user, 'ownerprofile'):
        return redirect('owner_login')

    # Fetch owner-related data
    properties = Property.objects.filter(
        owner=request.user
    ).order_by('-id')

    bookings = Booking.objects.filter(
        property__owner=request.user
    ).select_related(
        'property', 'user'
    ).order_by('-id')

    reviews = Review.objects.filter(
        property__owner=request.user
    ).select_related(
        'property', 'user'
    ).order_by('-id')

    # Create workbook
    workbook = Workbook()

    report_sheets = [
        ('Properties', properties),
        ('Bookings', bookings),
        ('Reviews', reviews),
    ]

    # =========================
    # SUMMARY SHEET
    # =========================

    summary = workbook.active
    summary.title = 'Summary'

    summary.append(['OWNER RENTAL REPORT'])

    summary.append([
        'Generated On',
        timezone.localtime().strftime('%d-%m-%Y %H:%M')
    ])

    summary.append([])
    summary.append(['Report Metric', 'Total'])
    summary.append(['Total Properties', properties.count()])
    summary.append(['Total Booking Requests', bookings.count()])
    summary.append(['Total Reviews', reviews.count()])

    summary['A1'].font = Font(bold=True, size=16)
    summary['A4'].font = Font(bold=True)
    summary['B4'].font = Font(bold=True)

    summary.column_dimensions['A'].width = 30
    summary.column_dimensions['B'].width = 25

    # =========================
    # CONVERT VALUES FOR EXCEL
    # =========================

    def convert_excel_value(value):

        if value is None:
            return ''

        # Convert lists, tuples, and sets
        if isinstance(value, (list, tuple, set)):
            return ', '.join(
                str(item) for item in value
            )

        # Convert dictionaries
        if isinstance(value, dict):
            return str(value)

        # Remove timezone information from datetime values
        if isinstance(value, datetime):

            if timezone.is_aware(value):
                value = timezone.make_naive(
                    value,
                    timezone.get_current_timezone()
                )

            return value

        # Remove timezone information from time values
        if isinstance(value, time):

            if value.tzinfo is not None:
                value = value.replace(tzinfo=None)

            return value

        # Keep supported Excel values unchanged
        if isinstance(value, (str, int, float, bool, date)):
            return value

        # Convert other unsupported values into text
        return str(value)

    # =========================
    # CREATE REPORT SHEETS
    # =========================

    for sheet_name, queryset in report_sheets:

        worksheet = workbook.create_sheet(
            title=sheet_name
        )

        # Get database model fields
        model_fields = [
            field
            for field in queryset.model._meta.concrete_fields
            if not field.many_to_many
        ]

        # Add headers
        headers = [
            field.verbose_name.title()
            for field in model_fields
        ]

        worksheet.append(headers)

        # Add database records
        for obj in queryset:

            row = []

            for field in model_fields:

                value = getattr(
                    obj,
                    field.attname
                )

                # Display readable foreign-key values
                if field.is_relation and field.many_to_one:

                    related_obj = getattr(
                        obj,
                        field.name
                    )

                    value = (
                        str(related_obj)
                        if related_obj is not None
                        else ''
                    )

                # Convert value before adding it to Excel
                value = convert_excel_value(value)

                row.append(value)

            worksheet.append(row)

        # =========================
        # STYLE HEADERS
        # =========================

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True,
                color='FFFFFF'
            )

            cell.fill = PatternFill(
                start_color='1F4E78',
                end_color='1F4E78',
                fill_type='solid'
            )

            cell.alignment = Alignment(
                horizontal='center',
                vertical='center'
            )

        # =========================
        # FORMAT COLUMNS
        # =========================

        worksheet.freeze_panes = 'A2'
        worksheet.auto_filter.ref = worksheet.dimensions

        for column_cells in worksheet.columns:

            max_length = 0

            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 3, 35)

    # =========================
    # DOWNLOAD EXCEL FILE
    # =========================

    response = HttpResponse(
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        )
    )

    response['Content-Disposition'] = (
        'attachment; filename="Owner_Rental_Reports.xlsx"'
    )

    workbook.save(response)

    return response

# =========================
# OWNER LOGOUT
# =========================
def owner_logout(request):

    logout(request)

    return redirect('owner_login')

from django.shortcuts import render

from .models import (
    HomeHero,
    HomeCategory,
    HomeArea,
    HomeFeature,
    HomeAISection,
    HomeCTA,
    HomeContact
)




def home(request):

    hero = HomeHero.objects.filter(
        is_active=True
    ).first()
    categories = HomeCategory.objects.filter(
        is_active=True
    ).order_by('display_order')
    why_choose_features = WhyChooseFeature.objects.filter(is_active=True).order_by('display_order')
    available_areas = AvailableArea.objects.filter(
    is_active=True
).order_by('display_order', 'name')
    return render(
        request,
        'home.html',
        {
            'hero': hero,
            'categories': categories,
            'why_choose_features': why_choose_features,
            'available_areas': available_areas
        }
    )

from .models import HomeHero
def homepage_management(request):

    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    heroes = HomeHero.objects.all().order_by('-created_at')
    categories = HomeCategory.objects.all().order_by('display_order')
    why_choose_features = WhyChooseFeature.objects.all()
    areas = AvailableArea.objects.all().order_by('display_order', 'name')
    return render(
        request,
        'admin_panel/homepage_management.html',
        {
            'heroes': heroes,
            'categories': categories,
            'why_choose_features': why_choose_features,
            'areas': areas
        }
    )
def add_hero(request):

    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    if request.method == 'POST':

        HomeHero.objects.create(
            badge=request.POST.get('badge'),
            title=request.POST.get('title'),
            description=request.POST.get('description'),
            button_text=request.POST.get('button_text'),
            button_link=request.POST.get('button_link'),
            image=request.FILES.get('image'),
            is_active=request.POST.get('is_active') == 'on'
        )

    return redirect('homepage_management')

def edit_hero(request, id):

    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    hero = get_object_or_404(HomeHero, id=id)

    if request.method == 'POST':

        hero.badge = request.POST.get('badge')
        hero.title = request.POST.get('title')
        hero.description = request.POST.get('description')
        hero.button_text = request.POST.get('button_text')
        hero.button_link = request.POST.get('button_link')

        # Update image only if a new image is selected
        if request.FILES.get('image'):
            hero.image = request.FILES.get('image')

        hero.is_active = request.POST.get('is_active') == 'on'

        hero.save()

    return redirect('homepage_management')



def delete_hero(request, id):

    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    hero = get_object_or_404(HomeHero, id=id)

    if request.method == 'POST':
        hero.delete()

    return redirect('homepage_management')



# Homepage Category Management Views
def add_category(request):
    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    if request.method == 'POST':
        name = request.POST.get('name', '').strip()

        if not name:
            messages.error(request, 'Category name is required.')
            return redirect('homepage_management')

        if HomeCategory.objects.filter(name__iexact=name).exists():
            messages.error(request, 'Category already exists.')
            return redirect('homepage_management')

        HomeCategory.objects.create(
            name=name,
            description=request.POST.get('description', '').strip(),
            icon=request.POST.get('icon', '').strip()
                 or 'fa-solid fa-house',
            display_order=request.POST.get('display_order') or 0,
            is_active=request.POST.get('is_active') == 'on'
        )

        messages.success(request, 'Category added successfully.')

    return redirect('homepage_management')

def edit_category(request, id):
    # Check admin login
    if not request.user.is_authenticated:
        return redirect('admin_login')

    # Allow only users who have an adminprofile
    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    category = get_object_or_404(HomeCategory, id=id)

    # Update only through POST
    if request.method != 'POST':
        return redirect('homepage_management')

    name = request.POST.get('name', '').strip()
    description = request.POST.get('description', '').strip()
    icon = request.POST.get('icon', '').strip()
    display_order = request.POST.get('display_order', '0')

    if not name:
        messages.error(request, 'Category name is required.')
        return redirect('homepage_management')

    if HomeCategory.objects.filter(
        name__iexact=name
    ).exclude(id=id).exists():
        messages.error(request, 'Category already exists.')
        return redirect('homepage_management')

    try:
        display_order = int(display_order)
        if display_order < 0:
            raise ValueError
    except (ValueError, TypeError):
        messages.error(request, 'Enter a valid display order.')
        return redirect('homepage_management')

    category.name = name
    category.description = description
    category.icon = icon or 'fa-solid fa-house'
    category.display_order = display_order
    category.is_active = request.POST.get('is_active') == 'on'

    category.save()

    messages.success(request, 'Category updated successfully.')
    return redirect('homepage_management')


def delete_category(request, id):
    # Check admin login
    if not request.user.is_authenticated:
        return redirect('admin_login')

    # Allow only users who have an adminprofile
    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    # Prevent deletion through a normal GET request
    if request.method != 'POST':
        return redirect('homepage_management')

    category = get_object_or_404(HomeCategory, id=id)
    category.delete()

    messages.success(request, 'Category deleted successfully.')
    return redirect('homepage_management')


from .models import WhyChooseFeature

def add_why_choose_feature(request):
    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    if request.method != 'POST':
        return redirect('homepage_management')

    title = request.POST.get('title', '').strip()
    description = request.POST.get('description', '').strip()
    icon = request.POST.get('icon', '').strip()
    display_order = request.POST.get('display_order', '0')

    if not title or not description:
        messages.error(request, 'Title and description are required.')
        return redirect('homepage_management')

    try:
        display_order = int(display_order)
        if display_order < 0:
            raise ValueError
    except (ValueError, TypeError):
        messages.error(request, 'Enter a valid display order.')
        return redirect('homepage_management')

    WhyChooseFeature.objects.create(
        title=title,
        description=description,
        icon=icon or 'fa-solid fa-house',
        display_order=display_order,
        is_active=request.POST.get('is_active') == 'on'
    )

    messages.success(request, 'Feature added successfully.')
    return redirect('homepage_management')


def edit_why_choose_feature(request, id):
    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    feature = get_object_or_404(WhyChooseFeature, id=id)

    if request.method != 'POST':
        return redirect('homepage_management')

    title = request.POST.get('title', '').strip()
    description = request.POST.get('description', '').strip()
    icon = request.POST.get('icon', '').strip()
    display_order = request.POST.get('display_order', '0')

    if not title or not description:
        messages.error(request, 'Title and description are required.')
        return redirect('homepage_management')

    try:
        display_order = int(display_order)
        if display_order < 0:
            raise ValueError
    except (ValueError, TypeError):
        messages.error(request, 'Enter a valid display order.')
        return redirect('homepage_management')

    feature.title = title
    feature.description = description
    feature.icon = icon or 'fa-solid fa-house'
    feature.display_order = display_order
    feature.is_active = request.POST.get('is_active') == 'on'
    feature.save()

    messages.success(request, 'Feature updated successfully.')
    return redirect('homepage_management')


def delete_why_choose_feature(request, id):
    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    if request.method != 'POST':
        return redirect('homepage_management')

    feature = get_object_or_404(WhyChooseFeature, id=id)
    feature.delete()

    messages.success(request, 'Feature deleted successfully.')
    return redirect('homepage_management')


# area management views
from .models import AvailableArea
# ================= ADD AVAILABLE AREA =================

def add_area(request):

    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    if request.method == 'POST':

        name = request.POST.get('name', '').strip()
        description = request.POST.get('description', '').strip()

        if name and description:

            if not AvailableArea.objects.filter(name__iexact=name).exists():

                AvailableArea.objects.create(
                    name=name,
                    description=description,
                    is_active=request.POST.get('is_active') == 'on'
                )

    return redirect('homepage_management')


# ================= EDIT AVAILABLE AREA =================

def edit_area(request, id):

    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    area = get_object_or_404(AvailableArea, id=id)

    if request.method == 'POST':

        name = request.POST.get('name', '').strip()
        description = request.POST.get('description', '').strip()

        if name and description:

            if not AvailableArea.objects.filter(
                name__iexact=name
            ).exclude(id=area.id).exists():

                area.name = name
                area.description = description
                area.is_active = request.POST.get('is_active') == 'on'

                area.save()

    return redirect('homepage_management')


# ================= DELETE AVAILABLE AREA =================

def delete_area(request, id):

    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    area = get_object_or_404(AvailableArea, id=id)

    if request.method == 'POST':
        area.delete()

    return redirect('homepage_management')



# Useer Management Views
# ================= USERS MANAGEMENT PAGE =================

from django.contrib.auth import get_user_model
from django.shortcuts import render, redirect


def users_management(request):

    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    User = get_user_model()

    users = User.objects.all().order_by('-date_joined')

    context = {
        'users': users,
        'total_users': users.count(),
    }

    return render(request, 'admin_panel/users_management.html', context)

# property management views
# ================= PROPERTIES MANAGEMENT =================
from .models import Property
def properties_management(request):

    if not request.user.is_authenticated:
        return redirect('admin_login')

    if not hasattr(request.user, 'adminprofile'):
        return redirect('admin_login')

    properties = Property.objects.all().order_by('-id')

    context = {
        'properties': properties,
        'total_properties': properties.count(),
    }

    return render(
        request,
        'admin_panel/properties_management.html',
        context
    )

# Property Management Views
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import Property, AvailableArea

# Owner Panel Views for Add Properties
from .models import Property, AvailableArea
from datetime import date
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import Property, AvailableArea


@login_required
def add_property(request):

    areas = AvailableArea.objects.filter(
        is_active=True
    ).order_by('display_order', 'name')

    if request.method == 'POST':

        category = request.POST.get('category')
        title = request.POST.get('title', '').strip()
        description = request.POST.get('description', '').strip()
        area = request.POST.get('area', '').strip()
        rent = request.POST.get('rent')

        furnishing = request.POST.get('furnishing', 'Unfurnished')
        amenities = request.POST.getlist('amenities')
        parking = request.POST.get('parking') == 'yes'
        bedrooms = request.POST.get('bedrooms') or 0
        bathrooms = request.POST.get('bathrooms') or 0
        floor_number = request.POST.get('floor_number') or 0
        available_from = request.POST.get('available_from') or None
        preferred_tenant = request.POST.get('preferred_tenant', 'Any')
        image = request.FILES.get('image')

        if not all([category, title, description, area, rent]):
            messages.error(request, 'Please fill in all required fields.')
            return render(
                request,
                'owner_panel/add_property.html',
                {'areas': areas}
            )

        if not areas.filter(name__iexact=area).exists():
            messages.error(request, 'Please select a valid area.')
            return render(
                request,
                'owner_panel/add_property.html',
                {'areas': areas}
            )

        try:
            rent = float(rent)

            if rent <= 0:
                raise ValueError

            bedrooms = int(bedrooms)
            bathrooms = int(bathrooms)
            floor_number = int(floor_number)

            if bedrooms < 0 or bathrooms < 0:
                raise ValueError

            if available_from:
                available_from = date.fromisoformat(available_from)

        except (ValueError, TypeError):
            messages.error(request, 'Please enter valid property details.')
            return render(
                request,
                'owner_panel/add_property.html',
                {'areas': areas}
            )

        Property.objects.create(
            owner=request.user,
            category=category,
            title=title,
            description=description,
            area=area,
            rent=rent,
            furnishing=furnishing,
            amenities=amenities,
            parking=parking,
            bedrooms=bedrooms,
            bathrooms=bathrooms,
            floor_number=floor_number,
            available_from=available_from,
            preferred_tenant=preferred_tenant,
            image=image
        )

        messages.success(request, 'Property added successfully!')
        return redirect('owner_dashboard')

    return render(
        request,
        'owner_panel/add_property.html',
        {'areas': areas}
    )
# Owner Panel Views for My Properties
@login_required
def my_properties(request):
    properties = Property.objects.filter(
        owner=request.user
    ).order_by('-created_at')

    return render(
        request,
        'owner_panel/my_properties.html',
        {'properties': properties}
    )
# myproperties edit and delete views
from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from .models import Property
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from .models import Property, AvailableArea


@login_required
def edit_property(request, property_id):

    # Get the property belonging to the logged-in owner
    property_obj = get_object_or_404(
        Property,
        id=property_id,
        owner=request.user
    )

    if request.method == 'POST':

        # Basic property details
        property_obj.category = request.POST.get('category')
        property_obj.title = request.POST.get('title', '').strip()
        property_obj.description = request.POST.get('description', '').strip()
        property_obj.area = request.POST.get('area', '').strip()
        property_obj.rent = request.POST.get('rent')

        # Furnishing
        property_obj.furnishing = request.POST.get('furnishing', '')

        # Amenities: multiple checkbox values
        property_obj.amenities = request.POST.getlist('amenities')

        # Parking
        property_obj.parking = (
            request.POST.get('parking') == 'yes'
        )

        # Bedrooms and bathrooms
        property_obj.bedrooms = request.POST.get('bedrooms') or None
        property_obj.bathrooms = request.POST.get('bathrooms') or None

        # Floor number
        property_obj.floor_number = (
            request.POST.get('floor_number') or None
        )

        # Available date
        property_obj.available_from = (
            request.POST.get('available_from') or None
        )

        # Preferred tenant
        property_obj.preferred_tenant = request.POST.get(
            'preferred_tenant', ''
        )

        # Replace image only when a new image is uploaded
        if request.FILES.get('image'):
            property_obj.image = request.FILES['image']

        property_obj.save()

        messages.success(
            request,
            'Property updated successfully!'
        )

        return redirect('owner_dashboard')

    context = {
        'property': property_obj,
        'available_areas': AvailableArea.objects.filter(
            is_active=True
        ).order_by('display_order', 'name'),

        'amenities_options': [
            'Wi-Fi',
            'Water Supply',
            'Power Backup',
            'Security',
            'Lift',
            'Washing Machine',
        ],
    }

    return render(
        request,
        'owner_panel/edit_property.html',
        context
    )

@login_required
def delete_property(request, property_id):
    if request.method == 'POST':
        property_obj = get_object_or_404(
            Property,
            id=property_id,
            owner=request.user
        )

        property_obj.delete()
        messages.success(request, 'Property deleted successfully.')

    return redirect('owner_dashboard')

# My Profile View

from django.contrib import messages
from django.contrib.auth import get_user_model, update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.db import IntegrityError
from django.shortcuts import redirect

User = get_user_model()


@login_required
def edit_owner_profile(request):
    if request.method != 'POST':
        return redirect('owner_dashboard')

    user = request.user

    # Allow only owners to edit their profiles
    if getattr(user, 'role', '') != 'OWNER':
        messages.error(request, 'Only owners can edit this profile.')
        return redirect('owner_dashboard')

    # Get form data
    username = request.POST.get('username', '').strip()
    email = request.POST.get('email', '').strip()
    new_password = request.POST.get('password', '')

    profile_url = '/owner-dashboard/#edit-profile-section'

    # Validate required fields
    if not username or not email:
        messages.error(request, 'Username and email are required.')
        return redirect(profile_url)

    # Check duplicate username
    if User.objects.filter(
        username=username
    ).exclude(pk=user.pk).exists():
        messages.error(request, 'This username is already taken.')
        return redirect(profile_url)

    # Check duplicate email
    if User.objects.filter(
        email=email
    ).exclude(pk=user.pk).exists():
        messages.error(request, 'This email is already in use.')
        return redirect(profile_url)

    # Update username and email
    user.username = username
    user.email = email

    # Update password only if a new password was entered
    if new_password:
        try:
            validate_password(new_password, user=user)
        except ValidationError as error:
            for error_message in error.messages:
                messages.error(request, error_message)
            return redirect(profile_url)

        user.set_password(new_password)

    # Save the changes
    try:
        user.save()
    except IntegrityError:
        messages.error(request, 'Unable to update profile. Please try again.')
        return redirect(profile_url)

    # Keep the owner logged in after changing the password
    if new_password:
        update_session_auth_hash(request, user)

    messages.success(request, 'Profile updated successfully!')
    return redirect('owner_dashboard')


# delete owner profile view


@login_required
def delete_owner_profile(request):
    if request.method != 'POST':
        return redirect('owner_dashboard')

    user = request.user

    if getattr(user, 'role', '') != 'OWNER':
        messages.error(request, 'Only owners can delete this profile.')
        return redirect('owner_dashboard')

    logout(request)
    user.delete()

    messages.success(request, 'Your profile has been deleted.')
    return redirect('owner_login')



# User Registration View

from django.shortcuts import render, redirect
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.urls import reverse
from .models import Property

def user_register(request):
    category = request.GET.get('category', '')

    if request.method == 'POST':
        category = request.POST.get('category', '')
        first_name = request.POST.get('first_name', '').strip()
        last_name = request.POST.get('last_name', '').strip()
        username = request.POST.get('username', '').strip()
        email = request.POST.get('email', '').strip()
        password = request.POST.get('password', '')
        confirm_password = request.POST.get('confirm_password', '')

        if password != confirm_password:
            messages.error(request, 'Passwords do not match.')
            return render(request, 'user_panel/user_register.html',
                          {'category': category})

        if User.objects.filter(username=username).exists():
            messages.error(request, 'Username already exists.')
            return render(request, 'user_panel/user_register.html',
                          {'category': category})

        if User.objects.filter(email=email).exists():
            messages.error(request, 'Email already registered.')
            return render(request, 'user_panel/user_register.html',
                          {'category': category})

        user = User.objects.create_user(
            username=username,
            email=email,
            password=password,
            first_name=first_name,
            last_name=last_name
        )

        # Remember the selected category until after login.
        request.session['selected_category'] = category

        messages.success(request, 'Registration successful. Please log in.')
        return redirect('user_login')

    return render(request, 'user_panel/user_register.html',
                  {'category': category})


def user_login(request):
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')
        category = request.POST.get('category', '')

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None:
            if not user.is_active:
                messages.error(
                    request,
                    'Your account is inactive.'
                )
                return redirect('user_login')

            login(request, user)

            # Preserve the selected category for the dashboard.
            if category:
                request.session['selected_category'] = category

            # Redirect to the logged-in user's dashboard.
            return redirect('user_dashboard')

        messages.error(
            request,
            'Invalid username or password.'
        )

    category = request.session.get('selected_category', '')

    return render(
        request,
        'user_panel/user_login.html',
        {'category': category}
    )


@login_required
def property_list(request):
    category = request.GET.get(
        'category',
        request.session.get('selected_category', '')
    )

    properties = Property.objects.filter(
        is_approved=True,
        is_available=True
    )

    if category:
        properties = properties.filter(property_type__iexact=category)

    return render(request, 'user_panel/user_properties.html', {
        'properties': properties,
        'category': category
    })


def user_logout(request):
    logout(request)
    return redirect('home')


# =========================
# USER DASHBOARD VIEW
# =========================

from datetime import date

from django.contrib import messages
from django.contrib.auth import (
    logout,
    update_session_auth_hash,
    get_user_model,
)
from django.contrib.auth.decorators import login_required
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.shortcuts import render, redirect
from django.urls import reverse

from .models import Property, Booking, Review

User = get_user_model()

from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.urls import reverse
from .models import Property


from datetime import date

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse

from .models import Property, Booking,Review

@login_required
def user_dashboard(request):

    # Prevent owners from accessing the User Dashboard
    if hasattr(request.user, 'ownerprofile'):
        return redirect('owner_dashboard')

    # Get the selected dashboard section
    section = request.GET.get('section', 'dashboard')

    # Fetch properties added by owners
    recent_properties = Property.objects.all().order_by('-created_at')

    # Current logged-in user
    user = request.user

    # Dashboard navigation URLs
    profile_url = reverse('user_dashboard') + '?section=profile'
    bookings_url = reverse('user_dashboard') + '?section=bookings'
    reviews_url = reverse('user_dashboard') + '?section=reviews'
    properties_url = reverse('user_dashboard') + '?section=properties'

    # Today's date
    today = date.today()

    # Get the selected property ID for the Reviews section
    property_id = request.GET.get('property_id')
    selected_property = None
    property_reviews = Review.objects.none()

    if property_id:
        selected_property = get_object_or_404(
            Property,
            pk=property_id
        )

    property_reviews = Review.objects.filter(
        property=selected_property
    ).select_related('user').order_by('-created_at')
    my_reviews = Review.objects.filter(
    user=request.user
    ).select_related('property').order_by('-created_at')
    # Allowed dashboard sections
    allowed_sections = [
        'dashboard',
        'profile',
        'properties',
        'bookings',
        'reviews',
        'recent_properties',
    ]

    if section not in allowed_sections:
        section = 'dashboard'

    # Handle POST requests
    if request.method == 'POST':

        action = request.POST.get('action')

        # Quick booking directly from a property card
        if action == 'quick_booking':

            property_id = request.POST.get('property_id')

            selected_property = get_object_or_404(
                Property,
                pk=property_id
            )

            # Prevent booking your own property
            if selected_property.owner_id == user.id:
                messages.error(
                    request,
                    "You cannot book your own property."
                )
                return redirect(properties_url)

            # Prevent duplicate pending booking requests
            existing_booking = Booking.objects.filter(
                user=user,
                property=selected_property,
                status='Pending'
            ).exists()

            if existing_booking:
                messages.warning(
                    request,
                    "You already have a pending request for this property."
                )
            else:
                Booking.objects.create(
                    user=user,
                    property=selected_property,
                    booking_date=today,
                    message="Quick booking request",
                    status='Pending'
                )

                messages.success(
                    request,
                    "Booking request submitted successfully!"
                )

            # Redirect to Booking History
            return redirect(bookings_url)

    # Fetch the logged-in user's booking requests
    my_bookings = Booking.objects.filter(
        user=user
    ).select_related('property').order_by('-pk')

    # Fetch the logged-in user's reviews
    my_reviews = Review.objects.filter(
        user=user
    ).select_related('property').order_by('-created_at')

    # Dashboard context
    context = {
        'user': user,
        'section': section,
        'recent_properties': recent_properties,
        'my_bookings': my_bookings,
        'my_reviews': my_reviews,
        'property_reviews': property_reviews,
        'selected_property': selected_property,
        'property_id': property_id,
        'today': today,
        'profile_url': profile_url,
        'bookings_url': bookings_url,
        'reviews_url': reviews_url,
        'properties_url': properties_url,
        'selected_property': selected_property,
        'property_reviews': property_reviews,
        'my_reviews': my_reviews,
    }

    if request.method == 'POST':
        action = request.POST.get('action')

        # =========================
        # EDIT PROFILE
        # =========================
        if action == 'edit_profile':

            username = request.POST.get('username', '').strip()
            first_name = request.POST.get('first_name', '').strip()
            last_name = request.POST.get('last_name', '').strip()
            email = request.POST.get('email', '').strip()

            if not username:
                messages.error(request, 'Username cannot be empty.')

            elif User.objects.filter(
                username=username
            ).exclude(pk=user.pk).exists():
                messages.error(request, 'This username is already taken.')

            elif email and User.objects.filter(
                email=email
            ).exclude(pk=user.pk).exists():
                messages.error(request, 'This email is already registered.')

            else:
                user.username = username
                user.first_name = first_name
                user.last_name = last_name
                user.email = email
                user.save()

                messages.success(
                    request,
                    'Profile updated successfully.'
                )

            return redirect(profile_url)

        # =========================
        # CHANGE PASSWORD
        # =========================
        elif action == 'change_password':

            old_password = request.POST.get('old_password', '')
            new_password = request.POST.get('new_password', '')
            confirm_password = request.POST.get('confirm_password', '')

            if not user.check_password(old_password):
                messages.error(
                    request,
                    'Your current password is incorrect.'
                )

            elif new_password != confirm_password:
                messages.error(
                    request,
                    'New passwords do not match.'
                )

            else:
                try:
                    validate_password(new_password, user=user)

                    user.set_password(new_password)
                    user.save()

                    update_session_auth_hash(request, user)

                    messages.success(
                        request,
                        'Password changed successfully.'
                    )

                except ValidationError as error:
                    for error_message in error.messages:
                        messages.error(request, error_message)

            return redirect(profile_url)

                # =========================
        # CREATE BOOKING REQUEST
        # =========================
        elif action == 'create_booking':

            property_id = request.POST.get('property_id')
            booking_date = request.POST.get('booking_date')
            booking_message = request.POST.get('message', '').strip()

            try:
                if not property_id or not booking_date:
                    raise ValueError

                property_obj = Property.objects.get(pk=property_id)
                visit_date = date.fromisoformat(booking_date)

                if visit_date < date.today():
                    messages.error(
                        request,
                        'Please select today or a future visit date.'
                    )
                    return redirect(bookings_url)

                Booking.objects.create(
                    user=request.user,
                    property=property_obj,
                    booking_date=visit_date,
                    message=booking_message,
                    status='Pending',
                )

                messages.success(
                    request,
                    'Your booking request was submitted successfully!'
                )

            except Property.DoesNotExist:
                messages.error(
                    request,
                    'The selected property does not exist.'
                )

            except (ValueError, TypeError):
                messages.error(
                    request,
                    'Please select a valid property and visit date.'
                )

            return redirect(bookings_url)


        # =========================
        # EDIT BOOKING REQUEST
        # =========================
        elif action == 'edit_booking':

            booking_id = request.POST.get('booking_id')
            property_id = request.POST.get('property_id')
            booking_date = request.POST.get('booking_date')
            booking_message = request.POST.get('message', '').strip()

            try:
                # Get only the booking belonging to the logged-in user
                booking = Booking.objects.get(
                    pk=booking_id,
                    user=request.user
                )

                if not property_id or not booking_date:
                    raise ValueError

                property_obj = Property.objects.get(pk=property_id)
                visit_date = date.fromisoformat(booking_date)

                if visit_date < date.today():
                    messages.error(
                        request,
                        'Please select today or a future visit date.'
                    )
                    return redirect(bookings_url)

                # Update booking details
                booking.property = property_obj
                booking.booking_date = visit_date
                booking.message = booking_message

                # Reset status after editing
                booking.status = 'Pending'

                booking.save()

                messages.success(
                    request,
                    'Booking request updated successfully!'
                )

            except Booking.DoesNotExist:
                messages.error(
                    request,
                    'Booking request not found or access denied.'
                )

            except Property.DoesNotExist:
                messages.error(
                    request,
                    'The selected property does not exist.'
                )

            except (ValueError, TypeError):
                messages.error(
                    request,
                    'Please select a valid property and visit date.'
                )

            return redirect(bookings_url)


        # =========================
        # DELETE BOOKING REQUEST
        # =========================
        elif action == 'delete_booking':

            booking_id = request.POST.get('booking_id')

            try:
                # Delete only the booking belonging to the logged-in user
                booking = Booking.objects.get(
                    pk=booking_id,
                    user=request.user
                )

                booking.delete()

                messages.success(
                    request,
                    'Booking request deleted successfully!'
                )

            except Booking.DoesNotExist:
                messages.error(
                    request,
                    'Booking request not found or access denied.'
                )

            return redirect(bookings_url)

                # =========================
        # SUBMIT REVIEW
        # =========================
        elif action == 'submit_review':

            property_id = request.POST.get('property_id')
            rating = request.POST.get('rating')
            comment = request.POST.get('comment', '').strip()

            reviews_url = reverse('user_dashboard') + '?section=reviews'

            try:
                if not property_id or not rating:
                    messages.error(
                        request,
                        'Please select a property and provide a rating.'
                    )
                    return redirect(reviews_url)

                property_obj = Property.objects.get(pk=property_id)
                rating = int(rating)

                if rating < 1 or rating > 5:
                    messages.error(
                        request,
                        'Please select a rating from 1 to 5.'
                    )
                    return redirect(reviews_url)

                Review.objects.create(
                    user=request.user,
                    property=property_obj,
                    rating=rating,
                    comment=comment
                )

                messages.success(
                    request,
                    'Your review was submitted successfully!'
                )

            except Property.DoesNotExist:
                messages.error(
                    request,
                    'The selected property does not exist.'
                )

            except (ValueError, TypeError):
                messages.error(
                    request,
                    'Please select a valid rating.'
                )

            return redirect(reviews_url)


        # =========================
        # EDIT REVIEW
        # =========================
        elif action == 'edit_review':

            review_id = request.POST.get('review_id')
            property_id = request.POST.get('property_id')
            rating = request.POST.get('rating')
            comment = request.POST.get('comment', '').strip()

            reviews_url = reverse('user_dashboard') + '?section=reviews'

            try:
                # Only allow editing the logged-in user's review
                review = Review.objects.get(
                    pk=review_id,
                    user=request.user
                )

                if not property_id or not rating:
                    messages.error(
                        request,
                        'Please select a property and provide a rating.'
                    )
                    return redirect(reviews_url)

                property_obj = Property.objects.get(pk=property_id)
                rating = int(rating)

                if rating < 1 or rating > 5:
                    messages.error(
                        request,
                        'Please select a rating from 1 to 5.'
                    )
                    return redirect(reviews_url)

                review.property = property_obj
                review.rating = rating
                review.comment = comment
                review.save()

                messages.success(
                    request,
                    'Your review was updated successfully!'
                )

            except Review.DoesNotExist:
                messages.error(
                    request,
                    'Review not found or access denied.'
                )

            except Property.DoesNotExist:
                messages.error(
                    request,
                    'The selected property does not exist.'
                )

            except (ValueError, TypeError):
                messages.error(
                    request,
                    'Please select a valid rating.'
                )

            return redirect(reviews_url)


        # =========================
        # DELETE REVIEW
        # =========================
        elif action == 'delete_review':

            review_id = request.POST.get('review_id')

            reviews_url = reverse('user_dashboard') + '?section=reviews'

            try:
                # Only allow deleting the logged-in user's review
                review = Review.objects.get(
                    pk=review_id,
                    user=request.user
                )

                review.delete()

                messages.success(
                    request,
                    'Your review was deleted successfully!'
                )

            except Review.DoesNotExist:
                messages.error(
                    request,
                    'Review not found or access denied.'
                )

            return redirect(reviews_url)
        # =========================
        # DELETE ACCOUNT
        # =========================
        elif action == 'delete_account':

            password = request.POST.get('delete_password', '')

            if not user.check_password(password):
                messages.error(
                    request,
                    'Incorrect password. Your account was not deleted.'
                )
                return redirect(profile_url)

            logout(request)
            user.delete()
            request.session.flush()

            return redirect('user_login')

    # =========================
    # LOAD PROPERTIES
    # =========================
    recent_properties = Property.objects.all().order_by('-id')

    # =========================
    # LOAD USER BOOKINGS
    # =========================
    my_bookings = Booking.objects.filter(
        user=user
    ).select_related(
        'property'
    ).order_by('-id')

    # =========================
    # LOAD USER REVIEWS
    # =========================
    my_reviews = Review.objects.filter(
        user=user
    ).select_related(
        'property'
    ).order_by('-id')

    # =========================
    # DASHBOARD CONTEXT
    # =========================
    context = {
        'section': section,
        'user': user,
        'total_properties': Property.objects.count(),
        'total_bookings': my_bookings.count(),
        'total_reviews': my_reviews.count(),
        'recent_properties': recent_properties,
        'my_bookings': my_bookings,
        'my_reviews': my_reviews,
    }

    return render(
        request,
        'user_panel/user_dashboard.html',
        context
    )



from django.contrib.auth.decorators import login_required
from django.shortcuts import render, redirect


@login_required
def user_profile(request):
    return render(
        request,
        'user_panel/user_profile.html',
        {'user': request.user}
    )

# Area viewa

from .models import AvailableArea
# propert details

# def admin_dashboard(request):

#     # Check whether the user is logged in
#     if not request.user.is_authenticated:
#         return redirect('admin_login')

#     # Check whether the logged-in user is an admin
#     if not AdminProfile.objects.filter(user=request.user).exists():
#         return redirect('admin_login')

#     # Count normal users (excluding admins and owners)
#     total_users = User.objects.filter(
#         adminprofile__isnull=True,
#         ownerprofile__isnull=True
#     ).count()

#     # Count registered owners
#     total_owners = OwnerProfile.objects.count()

#     # Send data to the dashboard template
#     context = {
#         'total_users': total_users,
#         'total_owners': total_owners,
#     }

#     return render(
#         request,
#         'admin_panel/admin_dashboard.html',
#         context
#     )
# Give me updated code